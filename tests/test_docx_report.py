from __future__ import annotations

from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.config import settings
from app.main import app
from app.reports.docx_report import DOCX_MIME_TYPE, markdown_to_docx_bytes
from app.reports.report_store import save_analysis_result


SAMPLE_MARKDOWN = """# ModernizeAI Modernization Report

> ModernizeAI performs static analysis only.

## 1. Executive Summary

This report summarizes modernization findings.

| Signal | Readout |
| --- | --- |
| Framework | Spring Boot |
| APIs | 5 |

- Add regression tests before migration.
"""


def test_markdown_to_docx_bytes_generates_readable_document() -> None:
    payload = markdown_to_docx_bytes(SAMPLE_MARKDOWN)

    assert payload.startswith(b"PK")
    assert len(payload) > 1000

    document = Document(BytesIO(payload))
    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "ModernizeAI Modernization Report" in paragraph_text
    assert "1. Executive Summary" in paragraph_text
    assert "This report summarizes modernization findings." in paragraph_text
    assert "Add regression tests before migration." in paragraph_text
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "Signal"
    assert document.tables[0].cell(1, 1).text == "Spring Boot"


def test_report_docx_endpoint_returns_word_download(tmp_path) -> None:
    original_report_dir = settings.report_dir
    object.__setattr__(settings, "report_dir", tmp_path)
    try:
        analysis_id = "docx-api-test"
        save_analysis_result(
            analysis_id,
            {
                "progress": [],
                "agent_graph": {},
                "agent_decisions": [],
                "final_report_markdown": SAMPLE_MARKDOWN,
            },
        )

        client = TestClient(app)
        markdown_response = client.get(f"/analysis/{analysis_id}/report")
        docx_response = client.get(f"/analysis/{analysis_id}/report.docx")
        missing_response = client.get("/analysis/missing-docx/report.docx")

        assert markdown_response.status_code == 200
        assert markdown_response.json()["report_markdown"] == SAMPLE_MARKDOWN
        assert docx_response.status_code == 200
        assert docx_response.headers["content-type"] == DOCX_MIME_TYPE
        assert (
            docx_response.headers["content-disposition"]
            == f'attachment; filename="modernizeai-report-{analysis_id}.docx"'
        )
        assert docx_response.content.startswith(b"PK")
        assert missing_response.status_code == 404
    finally:
        object.__setattr__(settings, "report_dir", original_report_dir)

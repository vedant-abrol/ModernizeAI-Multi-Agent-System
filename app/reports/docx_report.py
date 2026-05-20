from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.shared import Inches


DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def markdown_to_docx_bytes(markdown: str) -> bytes:
    document = Document()
    _set_margins(document)

    table_lines: list[str] = []
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if _is_table_line(line):
            table_lines.append(line)
            continue

        if table_lines:
            _add_table(document, table_lines)
            table_lines = []

        _add_line(document, line)

    if table_lines:
        _add_table(document, table_lines)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _set_margins(document: DocumentObject) -> None:
    for section in document.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)


def _add_line(document: DocumentObject, line: str) -> None:
    stripped = line.strip()
    if not stripped:
        return

    if stripped.startswith("# "):
        document.add_heading(_clean_inline(stripped[2:]), level=0)
        return
    if stripped.startswith("## "):
        document.add_heading(_clean_inline(stripped[3:]), level=1)
        return
    if stripped.startswith("### "):
        document.add_heading(_clean_inline(stripped[4:]), level=2)
        return
    if stripped.startswith("#### "):
        document.add_heading(_clean_inline(stripped[5:]), level=3)
        return

    if stripped.startswith(">"):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(_clean_inline(stripped.lstrip("> ")))
        run.italic = True
        return

    if stripped.startswith("- "):
        document.add_paragraph(_clean_inline(stripped[2:]), style="List Bullet")
        return

    document.add_paragraph(_clean_inline(stripped))


def _is_table_line(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|")


def _is_separator_row(cells: list[str]) -> bool:
    return all(cell and set(cell) <= {"-", ":"} for cell in cells)


def _parse_table_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [_clean_inline(cell.strip().replace("\\|", "|")) for cell in cells]


def _add_table(document: DocumentObject, table_lines: list[str]) -> None:
    rows = [_parse_table_row(line) for line in table_lines]
    rows = [row for row in rows if row and not _is_separator_row(row)]
    if not rows:
        return

    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"

    for row_index, row in enumerate(rows):
        for col_index in range(width):
            text = row[col_index] if col_index < len(row) else ""
            cell = table.cell(row_index, col_index)
            cell.text = text
            if row_index == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

    document.add_paragraph()


def _clean_inline(value: Any) -> str:
    text = str(value)
    replacements = {
        "**": "",
        "`": "",
        "\\|": "|",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text.strip()
